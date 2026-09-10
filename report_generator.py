# report_generator.py

import logging
from datetime import date
from typing import Optional

from colorama import init, Fore, Style
from tabulate import tabulate

from models import UserReport

from string import ascii_uppercase
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

init(autoreset=True)
logger = logging.getLogger(__name__)


class ReportGenerator:
    """Analiz sonuçlarını konsol ve Excel formatında raporlar."""

    DAY_NAMES = [
        "Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"
    ]

    def __init__(self, config: dict):
        self.config = config
        self.output_format = config["report"].get("output_format", "console")
        self.excel_file = config["report"].get(
            "excel_filename", "worklog_report.xlsx"
        )
        self.individual_sheets = config["report"].get(
            "individual_sheets", True
        )

    def generate(self, reports: list[UserReport]):
        """Tüm raporları üretir."""
        if self.output_format in ("console", "both"):
            self._print_console(reports)

        if self.output_format in ("excel", "both"):
            try:
                self._export_excel(reports)
            except Exception as e:
                logger.error(f"Excel raporu oluşturulurken hata oluştu: {e}")

    # ═══════════════════════════════════════════
    #  KONSOL RAPORU
    # ═══════════════════════════════════════════
    def _print_console(self, reports: list[UserReport]):
        print("\n")
        print(Fore.CYAN + "=" * 80)
        print(Fore.CYAN + "     JIRA WORKLOG KONTROL RAPORU")
        print(Fore.CYAN + "=" * 80)

        dr = self.config["date_range"]
        from datetime import datetime as _dt
        _start = _dt.strptime(str(dr['start_date']), "%d.%m.%Y")
        _end = _dt.strptime(str(dr['end_date']), "%d.%m.%Y")
        print(
            f"\n  📅 Tarih Aralığı: "
            f"{_start.strftime('%d.%m.%Y')} → {_end.strftime('%d.%m.%Y')}"
        )
        print(
            f"    Kontrol Edilen: "
            f"{len(reports)} çalışan"
        )
        filters = self.config["issue_filters"]
        print(
            f"    Filtreler: "
            f"tür={filters.get('issue_types')}, "
            f"statü={filters.get('statuses')}, "
            f"due_date={'zorunlu' if filters.get('require_due_date') else 'opsiyonel'}"
        )

        for report in reports:
            self._print_user_report(report)

        # Özet tablo
        self._print_summary_table(reports)

    def _print_user_report(self, report: UserReport):
        print(f"\n{'─' * 80}")
        print(
            Fore.YELLOW + Style.BRIGHT +
            f"\n    {report.display_name} ({report.username})"
        )
        print(f"{'─' * 80}")

        # ── YENİ: Başarı yüzdesi banner ──
        pct = report.success_percentage
        status = report.success_status
        pct_color = getattr(Fore, report.success_color)
        print(
            pct_color + Style.BRIGHT +
            f"  📊 Worklog Başarısı: %{pct} — {status}"
        )
        print(
            f"     (Worklog girilen gün: {report.total_days_worked} / "
            f"Toplam iş günü: {report.total_working_days})"
        )
        print(f"{'─' * 80}")

        # ── Uyarılar ──
        if report.warnings:
            print(Fore.RED + "\n    UYARILAR:")
            for w in report.warnings:
                print(Fore.RED + f"    {w}")

        if report.matching_issues:
            with_wl = [i for i in report.matching_issues if i["has_worklog"]]
            without_wl = [i for i in report.matching_issues if not i["has_worklog"]]

            def _fmt(seconds):
                """Saniyeyi okunabilir formata çevirir."""
                if seconds <= 0:
                    return "-"
                h = seconds // 3600
                m = (seconds % 3600) // 60
                return f"{h}s {m}dk" if m else f"{h}s"

            # Worklog'u OLAN issue'lar
            if with_wl:
                print(
                    Fore.GREEN +
                    f"\n  ✅ Worklog Girilen Issue'lar ({len(with_wl)}):"
                )
                issue_table = []
                for iss in with_wl:
                    prog = iss["workload_progress"]
                    prog_bar = self._progress_bar(prog)
                    issue_table.append([
                        iss["key"],
                        (iss["summary"][:30] + "..."
                         if len(iss["summary"]) > 30
                         else iss["summary"]),
                        iss["status"],
                        iss["due_date"],
                        _fmt(iss["original_estimate"]),
                        _fmt(iss["time_spent"]),
                        _fmt(iss["remaining_estimate"]),
                        f"{prog_bar} {prog}%",
                    ])
                print(tabulate(
                    issue_table,
                    headers=[
                        "Key", "Özet", "Statü", "Due Date",
                        "Tahmini", "Harcanan", "Kalan", "İlerleme"
                    ],
                    tablefmt="simple_outline",
                    stralign="left",
                ))

            # Worklog'u OLMAYAN issue'lar
            if without_wl:
                print(
                    Fore.RED +
                    f"\n  ❌ Worklog Girilmemiş Issue'lar ({len(without_wl)}):"
                )
                issue_table = []
                for iss in without_wl:
                    issue_table.append([
                        iss["key"],
                        (iss["summary"][:30] + "..."
                         if len(iss["summary"]) > 30
                         else iss["summary"]),
                        iss["status"],
                        iss["due_date"],
                        _fmt(iss["original_estimate"]),
                        _fmt(iss["time_spent"]),
                        _fmt(iss["remaining_estimate"]),
                    ])
                print(tabulate(
                    issue_table,
                    headers=[
                        "Key", "Özet", "Statü", "Due Date",
                        "Tahmini", "Harcanan", "Kalan"
                    ],
                    tablefmt="simple_outline",
                    stralign="left",
                ))

# ── YENİ: Workload Özeti ─���
            if report.total_original_estimate > 0:
                print(
                    Fore.MAGENTA +
                    f"\n  ⏱️  Workload Özeti:"
                )
                print(
                    f"     Toplam Tahmini : {_fmt(report.total_original_estimate)}"
                )
                print(
                    f"     Toplam Harcanan: {_fmt(report.total_logged_on_issues)}"
                )
                print(
                    f"     Toplam Kalan   : {_fmt(report.total_remaining_estimate)}"
                )
                overall_pct = round(
                    (report.total_logged_on_issues /
                     report.total_original_estimate) * 100, 1
                ) if report.total_original_estimate > 0 else 0
                print(
                    f"     Genel İlerleme : "
                    f"{self._progress_bar(overall_pct)} {overall_pct}%"
                )

        # ── Günlük Worklog Tablosu ──
        if report.daily_summaries:
            print(Fore.CYAN + "\n    Günlük Worklog Özeti:")

            table_data = []
            for d in sorted(report.daily_summaries.keys()):
                ds = report.daily_summaries[d]
                day_name = self.DAY_NAMES[d.weekday()]
                is_wknd = d.weekday() >= 5

                # Renklendirme bilgisi
                if is_wknd and ds.total_seconds == 0:
                    status = "─"
                elif ds.total_seconds == 0:
                    status = "  YOK"
                elif ds.total_hours < self.config["worklog_rules"].get(
                    "min_daily_hours", 0
                ):
                    status = "  AZ"
                elif ds.total_hours > self.config["worklog_rules"].get(
                    "max_daily_hours", 24
                ):
                    status = "  FAZLA"
                else:
                    status = "  OK"

                # Worklog detayları
                details = ""
                if ds.worklogs:
                    parts = []
                    for wl in ds.worklogs:
                        parts.append(
                            f"{wl.issue_key}({wl.time_spent_display})"
                        )
                    details = ", ".join(parts)

                row = [
                    d.strftime("%d.%m.%Y"),
                    day_name,
                    ds.total_display if ds.total_seconds > 0 else "-",
                    f"{ds.total_hours:.1f}" if ds.total_seconds > 0 else "-",
                    status,
                    (details[:50] + "..."
                     if len(details) > 50
                     else details),
                ]

                table_data.append(row)

            print(tabulate(
                table_data,
                headers=[
                    "Tarih", "Gün", "Süre",
                    "Saat", "Durum", "Detay"
                ],
                tablefmt="simple_outline",
                stralign="left",
                numalign="right",
            ))

             # Toplam satırının SONUNA başarı yüzdesini ekle:
            pct_color = getattr(Fore, report.success_color)
            print(
                Fore.WHITE + Style.BRIGHT +
                f"\n  📊 Toplam: {report.total_hours} saat | "
                f"Çalışılan gün: {report.total_days_worked}/{report.total_working_days} | "
                + pct_color + Style.BRIGHT +
                f"Başarı: %{report.success_percentage} {report.success_status}"
            )

  # ── YENİ: Progress bar yardımcı metodu ──
    @staticmethod
    def _progress_bar(percentage: float, width: int = 10) -> str:
        """Yüzdeye göre metin tabanlı ilerleme çubuğu."""
        pct = min(max(percentage, 0), 100)
        filled = int(width * pct / 100)
        empty = width - filled
        bar = "█" * filled + "░" * empty
        return f"[{bar}]"

        # _print_summary_table metodunu güncelle:

    def _print_summary_table(self, reports: list[UserReport]):
        print(f"\n{'═' * 90}")
        print(
            Fore.CYAN + Style.BRIGHT +
            "  📈 GENEL ÖZET"
        )
        print(f"{'═' * 90}")

        def _fmt(seconds):
            if seconds <= 0:
                return "-"
            h = seconds // 3600
            m = (seconds % 3600) // 60
            return f"{h}s" if not m else f"{h}s {m}dk"

        summary_data = []
        for r in reports:
            warning_count = len([
                w for w in r.warnings if "⚠️" in w or "❌" in w
            ])
            with_wl = len([i for i in r.matching_issues if i.get("has_worklog")])
            without_wl = len([i for i in r.matching_issues if not i.get("has_worklog")])
            summary_data.append([
                r.display_name,
                f"{with_wl}/{with_wl + without_wl}",
                f"{r.total_hours:.1f}",
                f"{r.total_days_worked}/{r.total_working_days}",
                _fmt(r.total_original_estimate),
                _fmt(r.total_remaining_estimate),
                f"%{r.success_percentage}",
                r.success_status,
            ])

        print(tabulate(
            summary_data,
            headers=[
                "Çalışan", "Issue(✅/Top)", "Worklog",
                "Gün(✅/Top)", "Tahmini", "Kalan",
                "Başarı", "Durum"
            ],
            tablefmt="simple_outline",
            stralign="left",
            numalign="right",
        ))
        print()

    # ═══════════════════════════════════════════
    #  EXCEL RAPORU
    # ═══════════════════════════════════════════
    # ── Stil sabitleri ──
    COLORS = {
        "primary":    "2F5496",
        "secondary":  "4472C4",
        "accent":     "5B9BD5",
        "success":    "548235",
        "warning":    "BF8F00",
        "alert":      "E36C09",
        "danger":     "C00000",
        "light_green":"C6EFCE",
        "light_yellow":"FFEB9C",
        "light_orange":"FBD5B5",
        "light_red":  "FFC7CE",
        "light_blue": "D6E4F0",
        "light_gray": "F2F2F2",
        "weekend":    "E2EFDA",
        "white":      "FFFFFF",
        "dark_text":  "1F3864",
    }

    def _get_styles(self):
        """Merkezi stil tanımları."""
        thin = Side(style="thin", color="B4C6E7")
        medium = Side(style="medium", color=self.COLORS["primary"])

        return {
            "title_font": Font(
                name="Segoe UI", bold=True, size=16,
                color=self.COLORS["white"]
            ),
            "subtitle_font": Font(
                name="Segoe UI", bold=True, size=11,
                color=self.COLORS["dark_text"]
            ),
            "header_font": Font(
                name="Segoe UI", bold=True, size=10,
                color=self.COLORS["white"]
            ),
            "data_font": Font(name="Segoe UI", size=10),
            "bold_font": Font(name="Segoe UI", bold=True, size=10),
            "small_font": Font(
                name="Segoe UI", size=9, color="666666"
            ),
            "header_fill": PatternFill(
                "solid", fgColor=self.COLORS["primary"]
            ),
            "subheader_fill": PatternFill(
                "solid", fgColor=self.COLORS["secondary"]
            ),
            "accent_fill": PatternFill(
                "solid", fgColor=self.COLORS["accent"]
            ),
            "success_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_green"]
            ),
            "warning_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_yellow"]
            ),
             "alert_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_orange"]
            ),
            "danger_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_red"]
            ),
            "stripe_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_blue"]
            ),
            "weekend_fill": PatternFill(
                "solid", fgColor=self.COLORS["weekend"]
            ),
            "light_gray_fill": PatternFill(
                "solid", fgColor=self.COLORS["light_gray"]
            ),
            "title_fill": PatternFill(
                "solid", fgColor=self.COLORS["primary"]
            ),
            "thin_border": Border(
                left=thin, right=thin, top=thin, bottom=thin
            ),
            "header_border": Border(
                left=thin, right=thin,
                top=medium, bottom=medium
            ),
            "center": Alignment(
                horizontal="center", vertical="center", wrap_text=True
            ),
            "left": Alignment(
                horizontal="left", vertical="center", wrap_text=True
            ),
            "right": Alignment(
                horizontal="right", vertical="center"
            ),
        }

    @staticmethod
    def _auto_fit_columns(ws, extra_padding: int = 3):
        """Sütun genişliklerini içeriğe göre otomatik ayarlar."""
        for col_cells in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column)

            for cell in col_cells:
                if cell.value is not None:
                    cell_text = str(cell.value)
                    # Çok satırlı hücrelerde en uzun satırı al
                    lines = cell_text.split("\n")
                    line_max = max(len(line) for line in lines)
                    # Bold fontlar biraz daha geniş
                    if cell.font and cell.font.bold:
                        line_max = int(line_max * 1.15)
                    max_length = max(max_length, line_max)

            # Min 8, Max 55 karakter genişlik
            adjusted = min(max(max_length + extra_padding, 8), 55)
            ws.column_dimensions[col_letter].width = adjusted

     # Sınıf değişkeni olarak ekle (class seviyesinde)
    _table_counter = 0

    @classmethod
    def _add_table(cls, ws, start_row, end_row, start_col, end_col, table_name):
        """Aralığa Excel tablosu ve filtre ekler."""
        import re
        import unicodedata

        # Veri satırı yoksa tablo ekleme (en az 1 başlık + 1 veri satırı)
        if end_row <= start_row:
            return

        # ── Tablo adı temizleme ──
        # 1. Türkçe karakterleri ASCII'ye çevir
        tr_map = str.maketrans(
            "çÇğĞıİöÖşŞüÜ",
            "cCgGiIoOsSuU"
        )
        safe_name = table_name.translate(tr_map)

        # 2. Unicode normalize
        safe_name = unicodedata.normalize("NFKD", safe_name)
        safe_name = safe_name.encode("ascii", "ignore").decode("ascii")

        # 3. Sadece harf, rakam ve _ bırak
        safe_name = re.sub(r"[^a-zA-Z0-9_]", "_", safe_name)

        # 4. Başta rakam varsa _ ekle
        if safe_name and safe_name[0].isdigit():
            safe_name = f"_{safe_name}"

        # 5. Boşsa varsayılan isim
        if not safe_name:
            safe_name = "Tablo"

        # 6. Benzersiz yap (global sayaç)
        cls._table_counter += 1
        safe_name = f"{safe_name[:200]}_{cls._table_counter}"

        # ── Tablo oluştur ──
        start_cell = f"{get_column_letter(start_col)}{start_row}"
        end_cell = f"{get_column_letter(end_col)}{end_row}"
        ref = f"{start_cell}:{end_cell}"

        table = Table(displayName=safe_name, ref=ref)
        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)

    @staticmethod
    def _fmt_seconds(seconds: int) -> str:
        if not seconds or seconds <= 0:
            return "-"
        h = seconds // 3600
        m = (seconds % 3600) // 60
        if h and m:
            return f"{h}s {m}dk"
        elif h:
            return f"{h}s"
        return f"{m}dk"

    def _export_excel(self, reports: list[UserReport]):
        """Profesyonel Excel raporu oluşturur."""
        ReportGenerator._table_counter = 0
        wb = Workbook()
        S = self._get_styles()

        # ══════════════════════════════════════
        #  SAYFA 1: DASHBOARD (Genel Özet)
        # ══════════════════════════════════════
        ws = wb.active
        ws.title = "Dashboard"
        ws.sheet_properties.tabColor = self.COLORS["primary"]

        # Başlık bandı
        ws.merge_cells("A1:I1")
        title_cell = ws.cell(row=1, column=1, value="📊 JIRA WORKLOG RAPORU")
        title_cell.font = S["title_font"]
        title_cell.fill = S["title_fill"]
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Meta bilgi
        dr = self.config["date_range"]
        filters = self.config["issue_filters"]
        from datetime import datetime as _dt
        _start = _dt.strptime(str(dr['start_date']), "%d.%m.%Y")
        _end = _dt.strptime(str(dr['end_date']), "%d.%m.%Y")
        meta_rows = [
            f"📅 Tarih Aralığı: {_start.strftime('%d.%m.%Y')} → {_end.strftime('%d.%m.%Y')}",
            f"👥 Çalışan Sayısı: {len(reports)}",
            f"🔍 Filtre: tür={filters.get('issue_types')}, statü={filters.get('statuses')}",
        ]
        for i, text in enumerate(meta_rows, 2):
            cell = ws.cell(row=i, column=1, value=text)
            cell.font = S["small_font"]
            ws.merge_cells(f"A{i}:I{i}")

        # ── Özet Tablosu ──
        summary_start = len(meta_rows) + 3
        headers = [
            "Çalışan", "Issue (✅/Top)", "Worklog (Saat)",
            "Gün (✅/Top)", "Tahmini", "Harcanan",
            "Kalan", "Başarı %", "Durum"
        ]

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=summary_start, column=col, value=h)
            cell.font = S["header_font"]
            cell.fill = S["header_fill"]
            cell.alignment = S["center"]
            cell.border = S["header_border"]

        for row_idx, r in enumerate(reports, summary_start + 1):
            with_wl = len([i for i in r.matching_issues if i.get("has_worklog")])
            total_iss = len(r.matching_issues)
            pct = r.success_percentage

            row_data = [
                r.display_name,
                f"{with_wl} / {total_iss}",
                r.total_hours,
                f"{r.total_days_worked} / {r.total_working_days}",
                self._fmt_seconds(r.total_original_estimate),
                self._fmt_seconds(r.total_logged_on_issues),
                self._fmt_seconds(r.total_remaining_estimate),
                pct,
                r.success_status,
            ]

            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.font = S["data_font"]
                cell.alignment = S["center"] if col > 1 else S["left"]
                cell.border = S["thin_border"]

            # Başarı hücresi renklendirme
            pct_cell = ws.cell(row=row_idx, column=8)
            status_cell = ws.cell(row=row_idx, column=9)
            if pct >= 90:
                pct_cell.fill = S["success_fill"]
                status_cell.fill = S["success_fill"]
                pct_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["success"]
                )
                status_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["success"]
                )
            elif pct >= 80:
                pct_cell.fill = S["warning_fill"]
                status_cell.fill = S["warning_fill"]
                pct_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["warning"]
                )
                status_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["warning"]
                )
            elif pct >= 70:
                pct_cell.fill = S["alert_fill"]
                status_cell.fill = S["alert_fill"]
                pct_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["alert"]
                )
                status_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["alert"]
                )
            else:
                pct_cell.fill = S["danger_fill"]
                status_cell.fill = S["danger_fill"]
                pct_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["danger"]
                )
                status_cell.font = Font(
                    name="Segoe UI", bold=True, size=10,
                    color=self.COLORS["danger"]
                )

        # Tablo + Filtre
        summary_end = summary_start + len(reports)
        if len(reports) > 0:
            self._add_table(
                ws, summary_start, summary_end,
                1, len(headers), "OzetTablosu"
            )

        # _export_excel içindeki grafik bölümünü komple SİL ve bununla değiştir:

        if len(reports) > 0:
            from openpyxl.chart import BarChart, Reference
            from openpyxl.chart.label import DataLabelList
            from openpyxl.chart.series import DataPoint
            from openpyxl.chart.shapes import GraphicalProperties
            from openpyxl.drawing.line import LineProperties

            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.width = 33
            chart.height = 14
            chart.title = "Başarı Oranları (%)"

            # Axis başlıkları (sadece metin, görsel yok)
            chart.y_axis.title = "Başarı %"
            chart.x_axis.title = "Çalışan"
            chart.y_axis.scaling.min = 0
            chart.y_axis.scaling.max = 105
            chart.y_axis.delete = False
            chart.x_axis.delete = False
            chart.y_axis.numFmt = '0"%"'

            # Legend kapalı (tek seri)
            chart.legend = None

            # Veri
            cats = Reference(
                ws, min_col=1,
                min_row=summary_start + 1,
                max_row=summary_end
            )
            vals = Reference(
                ws, min_col=8,
                min_row=summary_start,
                max_row=summary_end
            )
            chart.add_data(vals, titles_from_data=True)
            chart.set_categories(cats)

            # Her bar'ı başarıya göre renklendir
            COLOR_MAP = {
                "success": "548235",    # yeşil ≥90
                "warning": "BF8F00",    # sarı ≥80
                "alert":   "E36C09",    # turuncu ≥70 
                "danger":  "C00000",    # kırmızı <70
            }
            series = chart.series[0]
            for idx, r in enumerate(reports):
                pct = r.success_percentage
                color = (
                    COLOR_MAP["success"] if pct >= 90
                    else COLOR_MAP["warning"] if pct >= 80
                    else COLOR_MAP["alert"] if pct >= 70
                    else COLOR_MAP["danger"]
                )
                pt = DataPoint(idx=idx)
                pt.graphicalProperties = GraphicalProperties()
                pt.graphicalProperties.solidFill = color
                series.data_points.append(pt)

            # Data label
            #series.dLbls = DataLabelList()
            #series.dLbls.showVal = True
            #series.dLbls.numFmt = '0.0"%"'

            chart_row = summary_end + 2
            ws.add_chart(chart, f"A{chart_row}")

        # Freeze panes
        #ws.freeze_panes = f"A{summary_start + 1}"

        self._auto_fit_columns(ws)

        # ══════════════════════════════════════
        #  SAYFA 2+: Her kullanıcı ayrı sayfa
        # ══════════════════════════════════════
        if self.individual_sheets:
            for report in reports:
                safe_name = report.display_name[:28].replace("/", "-")
                ws_user = wb.create_sheet(title=safe_name)

                pct = report.success_percentage
                if pct >= 90:
                    ws_user.sheet_properties.tabColor = self.COLORS["success"]
                elif pct >= 80:
                    ws_user.sheet_properties.tabColor = self.COLORS["warning"]
                elif pct >= 70:
                    ws_user.sheet_properties.tabColor = self.COLORS["alert"]
                else:
                    ws_user.sheet_properties.tabColor = self.COLORS["danger"]

                current_row = 1

                # ── Kullanıcı Başlık Bandı ──
                ws_user.merge_cells(f"A{current_row}:H{current_row}")
                cell = ws_user.cell(
                    row=current_row, column=1,
                    value=f"👤 {report.display_name}"
                )
                cell.font = S["title_font"]
                cell.fill = S["title_fill"]
                cell.alignment = Alignment(horizontal="left", vertical="center")
                ws_user.row_dimensions[current_row].height = 36
                current_row += 1

                # ── Başarı Bandı ──
                ws_user.merge_cells(f"A{current_row}:H{current_row}")
                status_text = (
                    f"📊 Başarı: %{pct} — {report.success_status}  |  "
                    f"Çalışılan: {report.total_days_worked}/{report.total_working_days} gün  |  "
                    f"Toplam: {report.total_hours} saat"
                )
                cell = ws_user.cell(row=current_row, column=1, value=status_text)
                cell.font = Font(name="Segoe UI", bold=True, size=11)
                if pct >= 90:
                    cell.fill = S["success_fill"]
                elif pct >= 80:
                    cell.fill = S["warning_fill"]
                elif pct >= 70:
                    cell.fill = S["alert_fill"]
                else:
                    cell.fill = S["danger_fill"]
                cell.alignment = Alignment(horizontal="left", vertical="center")
                ws_user.row_dimensions[current_row].height = 28
                current_row += 2

                # ══════════════════════════
                #  BÖLÜM: Issue Tablosu
                # ══════════════════════════
                ws_user.cell(
                    row=current_row, column=1,
                    value="📋 Issue Listesi"
                ).font = S["subtitle_font"]
                current_row += 1

                issue_headers = [
                    "Key", "Özet", "Tür", "Statü", "Due Date",
                    "Tahmini", "Harcanan", "Kalan", "İlerleme %",
                    "Worklog Durumu"
                ]
                issue_header_row = current_row
                for col, h in enumerate(issue_headers, 1):
                    cell = ws_user.cell(row=current_row, column=col, value=h)
                    cell.font = S["header_font"]
                    cell.fill = S["header_fill"]
                    cell.alignment = S["center"]
                    cell.border = S["header_border"]
                current_row += 1

                for iss in report.matching_issues:
                    has_wl = iss.get("has_worklog", False)
                    prog = iss.get("workload_progress", 0)

                    row_data = [
                        iss["key"],
                        iss["summary"],
                        iss["type"],
                        iss["status"],
                        iss["due_date"],
                        self._fmt_seconds(iss.get("original_estimate", 0)),
                        self._fmt_seconds(iss.get("time_spent", 0)),
                        self._fmt_seconds(iss.get("remaining_estimate", 0)),
                        prog,
                        "✅ Girilmiş" if has_wl else "❌ Girilmemiş",
                    ]

                    for col, val in enumerate(row_data, 1):
                        cell = ws_user.cell(
                            row=current_row, column=col, value=val
                        )
                        cell.font = S["data_font"]
                        cell.border = S["thin_border"]
                        cell.alignment = S["center"] if col != 2 else S["left"]

                    # Satır renklendirme
                    wl_cell = ws_user.cell(row=current_row, column=10)
                    if has_wl:
                        wl_cell.fill = S["success_fill"]
                        wl_cell.font = Font(
                            name="Segoe UI", bold=True, size=10,
                            color=self.COLORS["success"]
                        )
                    else:
                        wl_cell.fill = S["danger_fill"]
                        wl_cell.font = Font(
                            name="Segoe UI", bold=True, size=10,
                            color=self.COLORS["danger"]
                        )

                    # İlerleme hücresi renk
                    prog_cell = ws_user.cell(row=current_row, column=9)
                    if prog >= 100:
                        prog_cell.fill = S["success_fill"]
                    elif prog >= 50:
                        prog_cell.fill = S["warning_fill"]
                    elif prog > 0:
                        prog_cell.fill = S["danger_fill"]

                    current_row += 1

                # Issue tablosu
                issue_end_row = current_row - 1
                if issue_end_row >= issue_header_row + 1:
                    self._add_table(
                        ws_user, issue_header_row, issue_end_row,
                        1, len(issue_headers),
                        f"Issues_{safe_name.replace(' ', '_')}"
                    )

                current_row += 1

                # ══════════════════════════════
                #  BÖLÜM: Workload Özeti
                # ══════════════════════════════
                if report.total_original_estimate > 0:
                    ws_user.cell(
                        row=current_row, column=1,
                        value="⏱️ Workload Özeti"
                    ).font = S["subtitle_font"]
                    current_row += 1

                    workload_items = [
                        ("Toplam Tahmini", self._fmt_seconds(report.total_original_estimate)),
                        ("Toplam Harcanan", self._fmt_seconds(report.total_logged_on_issues)),
                        ("Toplam Kalan", self._fmt_seconds(report.total_remaining_estimate)),
                        ("Genel İlerleme", f"%{round((report.total_logged_on_issues / max(report.total_original_estimate, 1)) * 100, 1)}"),
                    ]
                    for label, val in workload_items:
                        cell_l = ws_user.cell(
                            row=current_row, column=1, value=label
                        )
                        cell_l.font = S["bold_font"]
                        cell_l.fill = S["light_gray_fill"]
                        cell_l.border = S["thin_border"]

                        cell_v = ws_user.cell(
                            row=current_row, column=2, value=val
                        )
                        cell_v.font = S["data_font"]
                        cell_v.border = S["thin_border"]
                        cell_v.alignment = S["left"]
                        current_row += 1

                    current_row += 1

                # ══════════════════════════════════
                #  BÖLÜM: Günlük Worklog Detayları
                # ══════════════════════════════════
                ws_user.cell(
                    row=current_row, column=1,
                    value="📅 Günlük Worklog Detayları"
                ).font = S["subtitle_font"]
                current_row += 1

                day_headers = [
                    "Tarih", "Gün", "Toplam Saat", "Toplam Süre",
                    "Durum", "Issue Detayları", "Açıklamalar"
                ]
                day_header_row = current_row
                for col, h in enumerate(day_headers, 1):
                    cell = ws_user.cell(row=current_row, column=col, value=h)
                    cell.font = S["header_font"]
                    cell.fill = S["header_fill"]
                    cell.alignment = S["center"]
                    cell.border = S["header_border"]
                current_row += 1

                DAY_TR = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]
                min_h = self.config["worklog_rules"].get("min_daily_hours", 0)
                max_h = self.config["worklog_rules"].get("max_daily_hours", 24)

                for d in sorted(report.daily_summaries.keys()):
                    ds = report.daily_summaries[d]
                    day_name = DAY_TR[d.weekday()]
                    is_wknd = d.weekday() >= 5

                    if is_wknd and ds.total_seconds == 0:
                        status = "Hafta Sonu"
                    elif ds.total_seconds == 0:
                        status = "BOŞ"
                    elif ds.total_hours < min_h:
                        status = "Eksik Mesai"
                    elif ds.total_hours > max_h:
                        status = "Fazla Mesai"
                    else:
                        status = "OK"

                    details = "\n".join(
                        f"{wl.issue_key}: {wl.time_spent_display}"
                        for wl in ds.worklogs
                    )
                    comments = "\n".join(
                        f"{wl.issue_key}: {wl.comment[:40]}{'...' if len(wl.comment) > 40 else ''}"
                        for wl in ds.worklogs if wl.comment
                    )

                    row_data = [
                        d.strftime("%d.%m.%Y"),
                        day_name,
                        ds.total_hours if ds.total_seconds > 0 else 0,
                        ds.total_display if ds.total_seconds > 0 else "-",
                        status,
                        details or "-",
                        comments or "-",
                    ]

                    for col, val in enumerate(row_data, 1):
                        cell = ws_user.cell(
                            row=current_row, column=col, value=val
                        )
                        cell.font = S["data_font"]
                        cell.border = S["thin_border"]
                        cell.alignment = (
                            S["left"] if col >= 6 else S["center"]
                        )

                    # Satır renklendirme
                    fill = None
                    status_font = S["data_font"]
                    if is_wknd:
                        fill = S["weekend_fill"]
                    elif status == "BOŞ":
                        fill = S["danger_fill"]
                        status_font = Font(
                            name="Segoe UI", bold=True, size=10,
                            color=self.COLORS["danger"]
                        )
                    elif status == "Eksik Mesai":
                        fill = S["warning_fill"]
                        status_font = Font(
                            name="Segoe UI", bold=True, size=10,
                            color=self.COLORS["warning"]
                        )
                    elif status == "Fazla Mesai":
                        fill = S["stripe_fill"]
                        status_font = Font(
                            name="Segoe UI", bold=True, size=10,
                            color=self.COLORS["accent"]
                        )
                    elif status == "OK":
                        fill = S["success_fill"]

                    if fill:
                        for col in range(1, len(day_headers) + 1):
                            ws_user.cell(
                                row=current_row, column=col
                            ).fill = fill

                    ws_user.cell(
                        row=current_row, column=5
                    ).font = status_font

                    current_row += 1

                # Günlük tablo
                day_end_row = current_row - 1
                if day_end_row >= day_header_row + 1:
                    self._add_table(
                        ws_user, day_header_row, day_end_row,
                        1, len(day_headers),
                        f"Gunluk_{safe_name.replace(' ', '_')}"
                    )

                # ── Toplam satırı ──
                current_row += 1
                ws_user.merge_cells(
                    f"A{current_row}:B{current_row}"
                )
                cell = ws_user.cell(
                    row=current_row, column=1, value="TOPLAM"
                )
                cell.font = Font(
                    name="Segoe UI", bold=True, size=11,
                    color=self.COLORS["white"]
                )
                cell.fill = S["header_fill"]
                cell.alignment = S["center"]
                ws_user.cell(
                    row=current_row, column=2
                ).fill = S["header_fill"]

                cell = ws_user.cell(
                    row=current_row, column=3,
                    value=report.total_hours
                )
                cell.font = Font(
                    name="Segoe UI", bold=True, size=11,
                    color=self.COLORS["white"]
                )
                cell.fill = S["header_fill"]
                cell.alignment = S["center"]

                for col in range(4, len(day_headers) + 1):
                    ws_user.cell(
                        row=current_row, column=col
                    ).fill = S["header_fill"]

                current_row += 2

                # ══════════════════════════
                #  BÖLÜM: Uyarılar
                # ══════════════════════════
                if report.warnings:
                    ws_user.cell(
                        row=current_row, column=1,
                        value="🚨 Uyarılar"
                    ).font = Font(
                        name="Segoe UI", bold=True, size=12,
                        color=self.COLORS["danger"]
                    )
                    current_row += 1

                    warn_headers = ["#", "Uyarı Mesajı"]
                    warn_header_row = current_row
                    for col, h in enumerate(warn_headers, 1):
                        cell = ws_user.cell(
                            row=current_row, column=col, value=h
                        )
                        cell.font = S["header_font"]
                        cell.fill = PatternFill(
                            "solid", fgColor=self.COLORS["danger"]
                        )
                        cell.alignment = S["center"]
                        cell.border = S["header_border"]
                    current_row += 1

                    for idx, w in enumerate(report.warnings, 1):
                        ws_user.cell(
                            row=current_row, column=1, value=idx
                        ).font = S["data_font"]
                        ws_user.cell(
                            row=current_row, column=1
                        ).alignment = S["center"]
                        ws_user.cell(
                            row=current_row, column=1
                        ).border = S["thin_border"]

                        cell = ws_user.cell(
                            row=current_row, column=2, value=w
                        )
                        cell.font = S["data_font"]
                        cell.border = S["thin_border"]
                        cell.alignment = S["left"]

                        # Satır arka planı
                        if "❌" in w:
                            cell.fill = S["danger_fill"]
                        elif "⚠️" in w:
                            cell.fill = S["warning_fill"]

                        current_row += 1

                    warn_end_row = current_row - 1
                    if warn_end_row >= warn_header_row + 1:
                        self._add_table(
                            ws_user, warn_header_row, warn_end_row,
                            1, 2,
                            f"Uyarilar_{safe_name.replace(' ', '_')}"
                        )

                # Freeze panes — başlık bandının altı
                ws_user.freeze_panes = "A3"

                # Sütun genişliklerini otomatik ayarla
                self._auto_fit_columns(ws_user)

        # ══════════════════════════════════════
        #  SAYFA SON: Tüm Worklog'lar (Ham Veri)
        # ══════════════════════════════════════
        ws_raw = wb.create_sheet(title="Tüm Workloglar")
        ws_raw.sheet_properties.tabColor = "808080"

        raw_headers = [
            "Çalışan", "Tarih", "Gün", "Issue Key",
            "Issue Özet", "Issue Tür", "Issue Statü",
            "Süre (Saat)", "Süre", "Açıklama"
        ]
        for col, h in enumerate(raw_headers, 1):
            cell = ws_raw.cell(row=1, column=col, value=h)
            cell.font = S["header_font"]
            cell.fill = S["header_fill"]
            cell.alignment = S["center"]
            cell.border = S["header_border"]

        DAY_TR = ["Pzt", "Sal", "Çar", "Per", "Cum", "Cmt", "Paz"]
        raw_row = 2
        for report in reports:
            for d in sorted(report.daily_summaries.keys()):
                ds = report.daily_summaries[d]
                for wl in ds.worklogs:
                    row_data = [
                        report.display_name,
                        d.strftime("%d.%m.%Y"),
                        DAY_TR[d.weekday()],
                        wl.issue_key,
                        wl.issue_summary,
                        wl.issue_type,
                        wl.issue_status,
                        wl.hours,
                        wl.time_spent_display,
                        wl.comment or "-",
                    ]
                    for col, val in enumerate(row_data, 1):
                        cell = ws_raw.cell(
                            row=raw_row, column=col, value=val
                        )
                        cell.font = S["data_font"]
                        cell.border = S["thin_border"]
                        cell.alignment = (
                            S["left"] if col in (5, 10) else S["center"]
                        )
                    raw_row += 1

        if raw_row > 2:
            self._add_table(
                ws_raw, 1, raw_row - 1,
                1, len(raw_headers), "TumWorkloglar"
            )

        ws_raw.freeze_panes = "A2"
        self._auto_fit_columns(ws_raw)

        # ── Kaydet ──
        wb.save(self.excel_file)
        print(
            Fore.GREEN + Style.BRIGHT +
            f"\n  💾 Excel raporu kaydedildi: {self.excel_file}"
        )